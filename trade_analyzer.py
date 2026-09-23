import os
import sys
from datetime import datetime
import requests
import mysql.connector
from mysql.connector import Error
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================
# 1. DATABASE CONFIGURATION & DATA FETCHING
# ==========================================

def get_db_connection():
    """Database connection establish karta hai env variables se."""
    try:
        connection = mysql.connector.connect(
            host=os.getenv("DB_HOST", "localhost"),
            port=int(os.getenv("DB_PORT", 3306)),
            user=os.getenv("DB_USER", "root"),
            password=os.getenv("DB_PASS", ""),
            database=os.getenv("DB_NAME", "crypto_db"),
            connection_timeout=10
        )
        return connection
    except Error as e:
        print(f"[ERROR] Database connection failed: {e}")
        return None

def fetch_trades_from_db():
    """
    Database se trades fetch karta hai.
    Agar DB connect na ho to fallback sample data return karta hai.
    """
    connection = get_db_connection()
    trades = []

    if connection and connection.is_connected():
        try:
            cursor = connection.cursor(dictionary=True)
            # Apni DB table aur column names ke mutabiq query adjust karein
            query = """
                SELECT 
                    symbol,
                    start_time,
                    close_time,
                    close_reason,
                    max_floating_profit_pct,
                    net_profit,
                    tips
                FROM trades
                ORDER BY start_time DESC
            """
            cursor.execute(query)
            trades = cursor.fetchall()
            print(f"[INFO] Successfully fetched {len(trades)} trades from database.")
        except Error as e:
            print(f"[ERROR] Error executing query: {e}")
        finally:
            if connection.is_connected():
                cursor.close()
                connection.close()
    else:
        print("[WARNING] DB connection unavailable. Using sample trades for demonstration.")
        # Fallback Mock Data for Testing
        trades = [
            {
                "symbol": "BTCUSDT",
                "start_time": "2026-09-23 08:00:00",
                "close_time": "2026-09-23 12:30:00",
                "close_reason": "TP Hit",
                "max_floating_profit_pct": 0.052,  # 5.2%
                "net_profit": 320.50,
                "tips": "Clean breakout play. Trailing stop captured maximum move."
            },
            {
                "symbol": "ETHUSDT",
                "start_time": "2026-09-23 13:00:00",
                "close_time": "2026-09-23 15:45:00",
                "close_reason": "SL Hit",
                "max_floating_profit_pct": 0.024,  # 2.4% max up before dump
                "net_profit": -115.00,
                "tips": "Trade was +2.4% in profit before reversing. Shift SL to Breakeven after +2%."
            },
            {
                "symbol": "SOLUSDT",
                "start_time": "2026-09-23 16:10:00",
                "close_time": "2026-09-23 18:20:00",
                "close_reason": "TP Hit",
                "max_floating_profit_pct": 0.038,  # 3.8%
                "net_profit": 185.20,
                "tips": "Strong momentum entry at key support area."
            }
        ]

    return trades

# ==========================================
# 2. EXCEL REPORT GENERATION (OPENPYXL)
# ==========================================

def create_excel_report(trades_data, output_filename="Crypto_Trade_Report.xlsx"):
    """Trades data ko professional formatted Excel file mein save karta hai."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trade Analysis"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    HEADER_FILL = PatternFill(start_color="111827", end_color="111827", fill_type="solid") # Dark Charcoal
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    PROFIT_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # Light Green
    PROFIT_FONT = Font(name="Segoe UI", size=10, color="065F46", bold=True)
    
    LOSS_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light Red
    LOSS_FONT = Font(name="Segoe UI", size=10, color="991B1B", bold=True)
    
    REGULAR_FONT = Font(name="Segoe UI", size=10, color="1F2937")

    THIN_BORDER = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Table Headers
    headers = [
        "Pair",
        "Trade Start Time",
        "Trade Close Time",
        "Duration",
        "Close Reason",
        "Max Floating Profit (%)",
        "Net Profit / Loss ($)",
        "Trade Tips & Lessons"
    ]
    
    ws.append(headers)
    
    # Format Headers Row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28

    # Populate Trades Rows
    for row_idx, trade in enumerate(trades_data, start=2):
        # Calculate Duration
        start_str = str(trade["start_time"])
        close_str = str(trade["close_time"])
        
        try:
            start_dt = datetime.strptime(start_str, "%Y-%m-%d %H:%M:%S")
            close_dt = datetime.strptime(close_str, "%Y-%m-%d %H:%M:%S")
            duration_td = close_dt - start_dt
            hours, remainder = divmod(duration_td.total_seconds(), 3600)
            minutes, _ = divmod(remainder, 60)
            duration_str = f"{int(hours)}h {int(minutes)}m"
        except Exception:
            duration_str = "N/A"

        profit = float(trade.get("net_profit", 0.0))
        reason = str(trade.get("close_reason", "N/A"))
        max_float = float(trade.get("max_floating_profit_pct", 0.0))

        row_values = [
            trade.get("symbol", "N/A"),
            start_str,
            close_str,
            duration_str,
            reason,
            max_float,
            profit,
            trade.get("tips", "")
        ]
        
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 24

        # Apply Cell Borders and Alignments
        for col_num in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = THIN_BORDER
            cell.font = REGULAR_FONT
            cell.alignment = Alignment(vertical="center")

            # Center alignment for general columns
            if col_num in [1, 2, 3, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Column 5: Reason Formatting
            if col_num == 5:
                if "TP" in reason.upper():
                    cell.fill = PROFIT_FILL
                    cell.font = PROFIT_FONT
                elif "SL" in reason.upper():
                    cell.fill = LOSS_FILL
                    cell.font = LOSS_FONT

            # Column 6: Max Floating Profit (%)
            if col_num == 6:
                cell.number_format = '+0.00%;-0.00%;0.00%'
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # Column 7: Net Profit ($)
            if col_num == 7:
                cell.number_format = '$#,##0.00;($#,##0.00);"$0.00"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if profit > 0:
                    cell.fill = PROFIT_FILL
                    cell.font = PROFIT_FONT
                elif profit < 0:
                    cell.fill = LOSS_FILL
                    cell.font = LOSS_FONT

            # Column 8: Tips Wrap Text
            if col_num == 8:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Adjust Column Widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == 'H':
            ws.column_dimensions[col_letter].width = 45
        else:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_filename)
    print(f"[SUCCESS] Excel report saved to: {output_filename}")
    return output_filename

# ==========================================
# 3. NTFY NOTIFICATION ALERT
# ==========================================

def send_ntfy_notification(total_trades, net_pnl):
    """Execution update ntfy topic par push karta hai."""
    ntfy_topic = os.getenv("NTFY_TOPIC")
    if not ntfy_topic:
        print("[INFO] NTFY_TOPIC not configured. Skipping notification.")
        return

    pnl_symbol = "🟢 +" if net_pnl >= 0 else "🔴 "
    message = f"Analyzed {total_trades} trades.\nTotal PnL: {pnl_symbol}${net_pnl:.2f}\nExcel Report Generated Successfully."
    
    try:
        requests.post(
            f"https://ntfy.sh/{ntfy_topic}",
            data=message.encode('utf-8'),
            headers={
                "Title": "Crypto Trade Report Ready",
                "Priority": "default",
                "Tags": "chart_with_upwards_trend,file_folder"
            },
            timeout=10
        )
        print("[INFO] Notification sent to ntfy.")
    except Exception as e:
        print(f"[ERROR] Failed to send ntfy notification: {e}")

# ==========================================
# 4. MAIN RUNNER
# ==========================================

def main():
    print("--- Starting Trade Analysis Execution ---")
    trades = fetch_trades_from_db()
    
    if trades:
        report_file = create_excel_report(trades)
        total_pnl = sum(float(t.get("net_profit", 0.0)) for t in trades)
        send_ntfy_notification(len(trades), total_pnl)
    else:
        print("[WARNING] No trade data found to process.")

if __name__ == "__main__":
    main()
